from typing import List

from openpyxl import load_workbook

from proceso import Proceso


def _a_entero(valor, nombre, numero_fila):
    """Convierte una celda a entero o lanza ValueError indicando la fila."""
    if valor is None or str(valor).strip() == "":
        raise ValueError(f"Fila {numero_fila}: {nombre} está vacío.")
    if isinstance(valor, bool):
        raise ValueError(f"Fila {numero_fila}: {nombre} debe ser un número entero.")
    if isinstance(valor, (int, float)):
        if float(valor).is_integer():
            return int(valor)
        raise ValueError(f"Fila {numero_fila}: {nombre} debe ser un número entero.")
    try:
        return int(str(valor).strip())
    except ValueError:
        raise ValueError(f"Fila {numero_fila}: {nombre} debe ser un número entero.")


def _a_texto_id(valor):
    """El ID puede venir como texto (P1) o como número (1)."""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _es_encabezado(fila):
    """Si la columna de llegada no es numérica, la fila es un encabezado."""
    llegada = fila[1] if len(fila) > 1 else None
    if isinstance(llegada, (int, float)):
        return False
    try:
        int(str(llegada).strip())
        return False
    except ValueError:
        return True


def leer_procesos_excel(ruta: str) -> List[Proceso]:
    """Devuelve la lista de procesos del Excel o lanza ValueError con el motivo."""
    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except Exception:
        raise ValueError("No se pudo abrir el archivo. Verifique que sea un "
                         "Excel válido (.xlsx).")

    try:
        hoja = libro.worksheets[0]
        filas = [fila for fila in hoja.iter_rows(values_only=True)
                 if any(celda is not None and str(celda).strip() != ""
                        for celda in fila)]
    finally:
        libro.close()

    if not filas:
        raise ValueError("El archivo no contiene datos.")

    primera_fila_excel = 1
    if _es_encabezado(filas[0]):
        filas = filas[1:]
        primera_fila_excel = 2
    if not filas:
        raise ValueError("El archivo solo tiene encabezado, no hay procesos.")

    procesos = []
    ids_usados = set()

    for posicion, fila in enumerate(filas):
        numero_fila = primera_fila_excel + posicion
        if len(fila) < 4:
            raise ValueError(f"Fila {numero_fila}: faltan columnas "
                             "(se requieren ID, Llegada, Ejecución y Prioridad).")

        if fila[0] is None or str(fila[0]).strip() == "":
            raise ValueError(f"Fila {numero_fila}: el ID está vacío.")
        id_proceso = _a_texto_id(fila[0])
        if id_proceso.lower() in ids_usados:
            raise ValueError(f"Fila {numero_fila}: el ID '{id_proceso}' está repetido.")
        ids_usados.add(id_proceso.lower())

        llegada = _a_entero(fila[1], "El tiempo de llegada", numero_fila)
        if llegada < 0:
            raise ValueError(f"Fila {numero_fila}: la llegada no puede ser negativa.")

        ejecucion = _a_entero(fila[2], "El tiempo de ejecución", numero_fila)
        if ejecucion <= 0:
            raise ValueError(f"Fila {numero_fila}: la ejecución debe ser mayor que 0.")

        prioridad = _a_entero(fila[3], "La prioridad", numero_fila)
        if prioridad <= 0:
            raise ValueError(f"Fila {numero_fila}: la prioridad debe ser mayor que 0.")

        procesos.append(Proceso(id_proceso, llegada, ejecucion, prioridad))

    return procesos